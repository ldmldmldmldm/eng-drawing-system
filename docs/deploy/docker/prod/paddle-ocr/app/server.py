import base64
import io
import os
from dataclasses import dataclass
from datetime import datetime
from typing import List, Optional

from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from paddleocr import PaddleOCR


app = FastAPI(title="PaddleOCR Local Service")

OCR_LANG = os.getenv("OCR_LANG", "ch")
OCR_USE_GPU = os.getenv("OCR_USE_GPU", "false").lower() == "true"
LOGO_PATH = os.path.join(os.path.dirname(__file__), "assets", "logo.png")
LOW_SCORE_THRESHOLD = 0.6
FOOTER_BRAND = "冲压九厂AI"

# 品牌色（对齐 π logo 橙系）
COLOR_TITLE_BG = "FFF7ED"
COLOR_TITLE_TEXT = "9A3412"
COLOR_HEADER_BG = "FB923C"
COLOR_HEADER_TEXT = "FFFFFF"
COLOR_BORDER = "FED7AA"
COLOR_ZEBRA = "FFF7ED"
COLOR_LOW_SCORE = "FECACA"
COLOR_FOOTER_TEXT = "78716C"

ocr = PaddleOCR(
    use_angle_cls=True,
    lang=OCR_LANG,
    use_gpu=OCR_USE_GPU,
)


class OcrBase64Req(BaseModel):
    image_base64: str


class TableExcelBase64Req(BaseModel):
    image_base64: str
    title: Optional[str] = None


@dataclass
class CellData:
    text: str
    min_score: float = 1.0


@app.get("/health")
def health():
    return {"ok": True}


@app.post("/ocr/file")
async def ocr_file(file: UploadFile = File(...)):
    try:
        raw = await file.read()
        if not raw:
            raise HTTPException(status_code=400, detail="empty file")
        result = ocr.ocr(raw, cls=True)
        text = extract_text(result)
        return {"code": 0, "text": text}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ocr/base64")
def ocr_base64(req: OcrBase64Req):
    try:
        raw = base64.b64decode(req.image_base64)
        if not raw:
            raise HTTPException(status_code=400, detail="empty image")
        result = ocr.ocr(raw, cls=True)
        text = extract_text(result)
        return {"code": 0, "text": text}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ocr/table")
async def ocr_table(file: UploadFile = File(...)):
    """返回每个文字的 bounding box + 文本 + 置信度"""
    try:
        raw = await file.read()
        if not raw:
            raise HTTPException(status_code=400, detail="empty file")
        result = ocr.ocr(raw, cls=True)
        items = extract_items_with_boxes(result)
        return {"code": 0, "items": items, "count": len(items)}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/ocr/table-excel")
async def ocr_table_excel(
    file: UploadFile = File(...),
    title: str = "图片表格还原",
    plain: bool = True,
):
    """OCR 识别图片表格 → 坐标聚类还原行列 → 返回 Excel 文件

    plain=true（默认）：仅还原 OCR 表格，中性边框/字体，无 logo/品牌色/页脚。
    plain=false：第 1 行留空仅右上角 logo，表格从第 2 行起，页脚含冲压九厂AI。
    """
    try:
        raw = await file.read()
        if not raw:
            raise HTTPException(status_code=400, detail="empty file")
        result = ocr.ocr(raw, cls=True)
        items = extract_items_with_boxes(result)
        if not items:
            raise HTTPException(status_code=422, detail="未识别到任何文字")
        table = cluster_items_to_table(items)
        buf = build_excel(table, title, plain=plain)
        num_cols = max(len(r) for r in table) if table else 0
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={
                "Content-Disposition": 'attachment; filename="table.xlsx"',
                "X-Table-Rows": str(len(table)),
                "X-Table-Cols": str(num_cols),
            },
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------

def extract_text(result) -> str:
    lines = []
    if not result:
        return ""
    for page in result:
        if not page:
            continue
        for item in page:
            if len(item) >= 2 and item[1] and len(item[1]) >= 1:
                lines.append(item[1][0])
    return "\n".join(lines)


def extract_items_with_boxes(result) -> list:
    items = []
    if not result:
        return items
    for page in result:
        if not page:
            continue
        for item in page:
            if len(item) >= 2 and item[1] and len(item[1]) >= 1:
                box = item[0]
                text = str(item[1][0]).strip()
                if not text:
                    continue
                items.append({
                    "box": [[round(p[0], 1), round(p[1], 1)] for p in box],
                    "text": text,
                    "score": round(float(item[1][1]), 4),
                })
    return items


def _y_center(item: dict) -> float:
    ys = [p[1] for p in item["box"]]
    return (min(ys) + max(ys)) / 2


def _x_center(item: dict) -> float:
    xs = [p[0] for p in item["box"]]
    return (min(xs) + max(xs)) / 2


def _box_height(item: dict) -> float:
    ys = [p[1] for p in item["box"]]
    return max(ys) - min(ys)


def _box_width(item: dict) -> float:
    xs = [p[0] for p in item["box"]]
    return max(xs) - min(xs)


def _cluster_1d_centers(values: List[float], threshold: float) -> List[float]:
    if not values:
        return []
    sorted_vals = sorted(values)
    groups: List[List[float]] = [[sorted_vals[0]]]
    for v in sorted_vals[1:]:
        center = sum(groups[-1]) / len(groups[-1])
        if abs(v - center) <= threshold:
            groups[-1].append(v)
        else:
            groups.append([v])
    return [sum(g) / len(g) for g in groups]


def _nearest_index(value: float, centers: List[float]) -> int:
    if not centers:
        return 0
    best_i = 0
    best_d = abs(value - centers[0])
    for i, center in enumerate(centers[1:], start=1):
        d = abs(value - center)
        if d < best_d:
            best_d = d
            best_i = i
    return best_i


def cluster_items_to_table(items: list) -> List[List[CellData]]:
    """Y 聚类成行，全局 X 聚类成列，同格多框按 X 合并文本"""
    if not items:
        return []

    median_h = sorted(_box_height(it) for it in items)[len(items) // 2]
    row_threshold = max(median_h * 0.55, 8.0)

    median_w = sorted(_box_width(it) for it in items)[len(items) // 2]
    col_threshold = max(median_w * 0.65, 12.0)

    sorted_items = sorted(items, key=lambda it: _y_center(it))

    row_groups: list[list[dict]] = []
    current_row = [sorted_items[0]]
    current_y = _y_center(sorted_items[0])

    for it in sorted_items[1:]:
        yc = _y_center(it)
        if abs(yc - current_y) <= row_threshold:
            current_row.append(it)
            current_y = (current_y * (len(current_row) - 1) + yc) / len(current_row)
        else:
            row_groups.append(current_row)
            current_row = [it]
            current_y = yc
    row_groups.append(current_row)

    col_centers = _cluster_1d_centers([_x_center(it) for it in items], col_threshold)
    if not col_centers:
        return []

    num_cols = len(col_centers)
    table: List[List[CellData]] = []
    for rg in row_groups:
        bucket: List[List[dict]] = [[] for _ in range(num_cols)]
        for it in rg:
            ci = _nearest_index(_x_center(it), col_centers)
            bucket[ci].append(it)
        row: List[CellData] = []
        for cell_items in bucket:
            cell_items.sort(key=lambda it: _x_center(it))
            texts = [it["text"] for it in cell_items]
            min_score = min(it["score"] for it in cell_items) if cell_items else 1.0
            row.append(CellData(" ".join(texts).strip(), min_score))
        table.append(row)

    return table


def _auto_column_widths(ws, start_row: int, end_row: int, num_cols: int) -> None:
    from openpyxl.utils import get_column_letter

    for ci in range(1, num_cols + 1):
        max_len = 8
        for ri in range(start_row, end_row + 1):
            val = ws.cell(ri, ci).value
            if val is None:
                continue
            for line in str(val).splitlines():
                max_len = max(max_len, len(line))
        ws.column_dimensions[get_column_letter(ci)].width = min(max(max_len * 1.15 + 2, 10), 46)


def build_excel(
    table: List[List[CellData]],
    title: str = "图片表格还原",
    plain: bool = True,
) -> io.BytesIO:
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "表格还原"

    num_cols = max(len(r) for r in table) if table else 1
    data_start_row = 1 if plain else 2
    data_end_row = data_start_row + len(table) - 1 if table else data_start_row

    if plain:
        plain_border = Border(
            left=Side(style="thin", color="D4D4D4"),
            right=Side(style="thin", color="D4D4D4"),
            top=Side(style="thin", color="D4D4D4"),
            bottom=Side(style="thin", color="D4D4D4"),
        )
        cell_font = Font(size=10, name="微软雅黑", color="000000")
        cell_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
        low_score_fill = PatternFill(start_color="F5F5F5", end_color="F5F5F5", fill_type="solid")

        for ri, row in enumerate(table, start=1):
            for ci, cell in enumerate(row, start=1):
                c = ws.cell(ri, ci, cell.text)
                c.font = cell_font
                c.alignment = cell_align
                c.border = plain_border
                if cell.min_score < LOW_SCORE_THRESHOLD:
                    c.fill = low_score_fill
            ws.row_dimensions[ri].height = 20
    else:
        thin = Border(
            left=Side(style="thin", color=COLOR_BORDER),
            right=Side(style="thin", color=COLOR_BORDER),
            top=Side(style="thin", color=COLOR_BORDER),
            bottom=Side(style="thin", color=COLOR_BORDER),
        )
        header_font = Font(bold=True, size=11, name="微软雅黑", color=COLOR_HEADER_TEXT)
        data_font = Font(size=10, name="微软雅黑", color="1C1917")
        footer_font = Font(size=9, name="微软雅黑", color=COLOR_FOOTER_TEXT)
        center = Alignment(horizontal="center", vertical="center", wrap_text=True)
        left = Alignment(horizontal="left", vertical="center", wrap_text=True)

        title_fill = PatternFill(start_color=COLOR_TITLE_BG, end_color=COLOR_TITLE_BG, fill_type="solid")
        header_fill = PatternFill(start_color=COLOR_HEADER_BG, end_color=COLOR_HEADER_BG, fill_type="solid")
        zebra_fill = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")
        low_score_fill = PatternFill(start_color=COLOR_LOW_SCORE, end_color=COLOR_LOW_SCORE, fill_type="solid")

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
        title_cell = ws.cell(1, 1, None)
        title_cell.alignment = center
        title_cell.fill = title_fill
        title_cell.border = thin
        for ci in range(2, num_cols + 1):
            c = ws.cell(1, ci)
            c.border = thin
            c.fill = title_fill
        ws.row_dimensions[1].height = 46

        if os.path.isfile(LOGO_PATH):
            try:
                from openpyxl.drawing.image import Image as XLImage

                logo = XLImage(LOGO_PATH)
                logo.height = 34
                logo.width = 34
                anchor_col = get_column_letter(num_cols)
                ws.add_image(logo, f"{anchor_col}1")
            except Exception:
                pass

        for ri, row in enumerate(table, start=data_start_row):
            is_header = ri == data_start_row
            is_zebra = (ri - data_start_row) % 2 == 1
            for ci, cell in enumerate(row, start=1):
                val = cell.text
                c = ws.cell(ri, ci, val)
                c.font = header_font if is_header else data_font
                c.alignment = center if is_header else left
                c.border = thin
                if is_header:
                    c.fill = header_fill
                elif cell.min_score < LOW_SCORE_THRESHOLD:
                    c.fill = low_score_fill
                elif is_zebra:
                    c.fill = zebra_fill
            ws.row_dimensions[ri].height = 24 if is_header else 22

        footer_row = data_end_row + 1
        ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=num_cols)
        footer_text = (
            f"{FOOTER_BRAND} · 图片表格还原  |  生成时间 {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  "
            f"共 {len(table)} 行 × {num_cols} 列"
        )
        footer_cell = ws.cell(footer_row, 1, footer_text)
        footer_cell.font = footer_font
        footer_cell.alignment = Alignment(horizontal="right", vertical="center")
        ws.row_dimensions[footer_row].height = 20

    _auto_column_widths(ws, data_start_row, data_end_row, num_cols)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
